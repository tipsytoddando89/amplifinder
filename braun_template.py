from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# Page margins
section = doc.sections[0]
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.85)

def set_para_shading(para, fill_hex):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    pPr.append(shd)

def set_para_border(para, sides, color_hex, size=4):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    for side in sides:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:space'), '1')
        el.set(qn('w:color'), color_hex)
        pBdr.append(el)
    pPr.append(pBdr)

def set_indent(para, left=0, right=0, hanging=0):
    pPr = para._p.get_or_add_pPr()
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), str(left))
    if right: ind.set(qn('w:right'), str(right))
    if hanging: ind.set(qn('w:hanging'), str(hanging))
    pPr.append(ind)

def set_spacing(para, before=0, after=0):
    pPr = para._p.get_or_add_pPr()
    spc = OxmlElement('w:spacing')
    spc.set(qn('w:before'), str(before))
    spc.set(qn('w:after'), str(after))
    pPr.append(spc)

# ── Title ──────────────────────────────────────────────
p = doc.add_paragraph()
set_spacing(p, 0, 40)
run = p.add_run("Braun Smart Home")
run.bold = True
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
run.font.name = "Arial"

p2 = doc.add_paragraph()
set_spacing(p2, 0, 160)
set_para_border(p2, ['bottom'], 'C9A84C', size=12)
run2 = p2.add_run("Discussion Questions & Response Template")
run2.font.size = Pt(12)
run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
run2.font.name = "Arial"

def section_header(num, title):
    p = doc.add_paragraph()
    set_spacing(p, 200, 80)
    set_para_shading(p, '1F3864')
    set_indent(p, left=180, right=180)
    run = p.add_run(f"  {num}. {title}")
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.size = Pt(14)
    run.font.name = "Arial"

def sub_question(text):
    p = doc.add_paragraph()
    set_spacing(p, 120, 40)
    set_indent(p, left=360)
    run = p.add_run(f"\u25b6  {text}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    run.font.name = "Arial"

def note_line(text):
    p = doc.add_paragraph()
    set_spacing(p, 40, 0)
    set_indent(p, left=720)
    run = p.add_run(f"\u2192 {text}")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    run.font.name = "Arial"

def bullet_line(label, text=""):
    p = doc.add_paragraph()
    set_spacing(p, 40, 0)
    set_indent(p, left=720, hanging=200)
    r1 = p.add_run("\u2022  ")
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    if label:
        r2 = p.add_run(label)
        r2.bold = True
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        r2.font.name = "Arial"
    if text:
        r3 = p.add_run(text)
        r3.font.size = Pt(10)
        r3.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        r3.font.name = "Arial"

def response_box():
    sides_all = ['top', 'bottom', 'left', 'right']
    sides_no_top = ['bottom', 'left', 'right']
    for i in range(4):
        p = doc.add_paragraph()
        set_para_shading(p, 'F5F5F5')
        set_indent(p, left=200, right=200)
        sides = sides_all if i == 0 else sides_no_top
        set_para_border(p, sides, 'CCCCCC', size=4)
        sp_before = 100 if i == 0 else 0
        sp_after = 160 if i == 3 else 0
        set_spacing(p, sp_before, sp_after)
        if i == 0:
            run = p.add_run("Response: ")
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            run.font.name = "Arial"
        else:
            run = p.add_run(" ")
            run.font.size = Pt(10)

# ── Section 1 ──────────────────────────────────────────
section_header("1", "Distribution Model (must include B2C)")

sub_question("Which retailers / integrators are envisioned?")
bullet_line("Custom Integrators: ", "See list attached")
bullet_line("Retail Customers: ", "")
response_box()

sub_question("Market sequencing: US first, then Europe?")
note_line("US & China at the same time")
note_line("Then Europe")
response_box()

# ── Section 2 ──────────────────────────────────────────
section_header("2", "Pricing & Value Proposition")

sub_question("Current pricing (2\u20135x Aqara benchmarks) is very high. What is the clear consumer POD justifying this premium vs. other smart home players?")
response_box()

out = "/Users/toddanderson/Documents/claude projects/amplifinder/Braun_SmartHome_Questions.docx"
doc.save(out)
print(f"Saved: {out}")
